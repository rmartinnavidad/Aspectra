from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

src = r'C:\Users\rmart\Documents\Codex\2026-09-06\ki\outputs\Aspectra\Aspectra Current Feature Reference.md'
out = r'C:\Users\rmart\Documents\Codex\2026-09-06\ki\outputs\Aspectra\Aspectra Current Feature Reference.docx'
doc=Document()
sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
styles=doc.styles
styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10.5)
styles['Title'].font.name='Aptos Display';styles['Title'].font.size=Pt(24);styles['Title'].font.bold=True
for h in ['Heading 1','Heading 2']:
    styles[h].font.name='Aptos Display';styles[h].font.color.rgb=None
title=doc.add_paragraph(style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER;title.add_run('Aspectra Current Feature Reference')
sub=doc.add_paragraph();sub.alignment=WD_ALIGN_PARAGRAPH.CENTER;sub.add_run('Current shipped behavior and known implementation boundaries').italic=True
for raw in open(src,encoding='utf-8'):
    line=raw.rstrip()
    if not line or line.startswith('# Aspectra'): continue
    if line.startswith('## '): doc.add_heading(line[3:],level=1)
    elif line.startswith('- '): doc.add_paragraph(line[2:],style='List Bullet')
    else: doc.add_paragraph(line)
doc.save(out)
